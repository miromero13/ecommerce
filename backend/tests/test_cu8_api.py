"""CU8 HTTP checks on an isolated in-memory database; never read dotenv."""
import importlib
import os
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from uuid import uuid4

import pytest
from openpyxl import load_workbook


@pytest.fixture
def cu8(monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from sqlalchemy.pool import StaticPool

    # No application secrets, production engine connection, or production DB schema.
    with patch.dict(os.environ, {'DATABASE_URL': 'sqlite://', 'SECRET_KEY': 'cu8-tests-only',
                                'ACCESS_TOKEN_EXPIRE_MINUTES': '60'}, clear=True), patch(
        'pydantic_settings.sources.DotEnvSettingsSource._read_env_files', return_value={}
    ):
        importlib.import_module('app.core.config')
        for path in Path(__file__).parents[1].joinpath('app/models').glob('*.py'):
            importlib.import_module('app.models.' + path.stem)
        from app.core.database import Base, get_db
        from app.auth.dependencies import get_current_payload
        from app.routes import report_routes
    engine = create_engine('sqlite://', poolclass=StaticPool, connect_args={'check_same_thread': False})
    tables = ['tenants', 'branches', 'products', 'product_variants', 'inventory', 'sales', 'sale_items', 'orders', 'order_items', 'inventory_movements']
    Base.metadata.create_all(engine, tables=[Base.metadata.tables[name] for name in tables])
    db = Session(engine)
    from app.models.tenant import Tenant
    from app.models.branch import Branch
    from app.models.product import Product
    from app.models.product_variant import ProductVariant
    from app.models.inventory import Inventory
    from app.models.sale import Sale
    from app.models.sale_item import SaleItem
    from app.models.order import Order
    from app.models.order_item import OrderItem

    companies = [Tenant(name=name, slug=name.lower(), contact_email=name + '@example.com') for name in ['AlphaCU8', 'BetaCU8']]
    db.add_all(companies)
    db.flush()
    branches, products = [], []
    cashier = uuid4()
    customer = uuid4()
    for index, company in enumerate(companies):
        branch = Branch(tenant_id=company.id, name=company.name, city='La Paz')
        product = Product(tenant_id=company.id, name=company.name + ' Blusa', category_id=uuid4())
        db.add_all([branch, product])
        db.flush()
        variant = ProductVariant(tenant_id=company.id, product_id=product.id, sku=company.slug, price=25)
        db.add(variant)
        db.flush()
        db.add(Inventory(tenant_id=company.id, branch_id=branch.id, variant_id=variant.id,
                         quantity=8 if index == 0 else 100, reserved_quantity=4, minimum_stock=5))
        for status, quantity in [('completed', 2 if index == 0 else 90), ('cancelled', 20)]:
            sale = Sale(tenant_id=company.id, branch_id=branch.id, user_id=cashier, status=status,
                        payment_method='cash', payment_status='paid', total_amount=25 * quantity,
                        created_at=datetime.now() - timedelta(days=14))
            db.add(sale)
            db.flush()
            db.add(SaleItem(sale_id=sale.id, variant_id=variant.id, product_id=product.id,
                            product_name=product.name, variant_sku=variant.sku, quantity=quantity,
                            unit_price=25, line_total=25 * quantity))
        order = Order(tenant_id=company.id, pickup_branch_id=branch.id, user_id=customer,
                      status='paid', payment_status='paid', payment_method='stripe', total_amount=25,
                      created_at=datetime.now() - timedelta(days=14))
        db.add(order)
        db.flush()
        db.add(OrderItem(order_id=order.id, variant_id=variant.id, product_id=product.id,
                         product_name=product.name, variant_sku=variant.sku, quantity=1, unit_price=25, line_total=25))
        branches.append(branch)
        products.append(product)
    db.commit()
    payload = {'sub': str(cashier), 'rol': 'administrador', 'tenant_id': str(companies[0].id), 'branch_id': None}
    app = FastAPI()
    app.include_router(report_routes.router, prefix='/api')
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_payload] = lambda: payload
    enabled = {'ml': True}
    monkeypatch.setattr(report_routes, 'tenant_has_feature', lambda db, tenant, key: enabled['ml'])
    with TestClient(app) as client:
        yield SimpleNamespace(client=client, db=db, payload=payload, companies=companies, branches=branches,
                              products=products, enabled=enabled)
    db.close()
    engine.dispose()


@pytest.mark.parametrize('path', ['top-products', 'branch-comparison', 'critical-inventory'])
def test_descriptive_reports_are_real_and_tenant_scoped(cu8, path):
    result = cu8.client.get('/api/reports/' + path)
    assert result.status_code == 200, result.text
    rows = result.json()['data']['rows']
    assert len(rows) == 1
    assert 'BetaCU8' not in result.text
    if path == 'critical-inventory':
        assert rows[0]['available_quantity'] == 4
        assert rows[0]['minimum_stock'] == 5
    else:
        assert rows[0]['quantity_sold'] == 3  # Two completed POS units + one paid web unit, not cancelled.
        assert float(rows[0]['gross_sales']) == 75
    foreign = cu8.client.get('/api/reports/' + path, params={'branch_id': str(cu8.branches[1].id)})
    assert foreign.status_code == 200
    assert foreign.json()['data']['rows'] == []


@pytest.mark.parametrize('path', ['sales', 'inventory', 'movements'])
def test_previous_json_and_csv_remain_available(cu8, path):
    result = cu8.client.get('/api/reports/' + path)
    assert result.status_code == 200, result.text
    assert 'BetaCU8' not in result.text
    csv = cu8.client.get('/api/reports/' + path + '/export')
    assert csv.status_code == 200
    assert 'text/csv' in csv.headers['content-type']
    assert 'BetaCU8' not in csv.text


@pytest.mark.parametrize('kind', ['sales', 'inventory', 'movements', 'top_products', 'branch_comparison', 'critical_inventory'])
@pytest.mark.parametrize('fmt', ['excel', 'pdf'])
def test_export_all_report_types(cu8, kind, fmt):
    result = cu8.client.get('/api/reports/export/' + fmt, params={'type': kind})
    assert result.status_code == 200, result.text
    if fmt == 'excel':
        rows = list(load_workbook(BytesIO(result.content)).active.values)
        assert 'BetaCU8' not in str(rows)
        assert result.content.startswith(b'PK')
    else:
        assert result.content.startswith(b'%PDF-')
    assert 'attachment;' in result.headers['content-disposition']


def test_ml_feature_role_and_branch_boundaries_leave_business_data_unchanged(cu8):
    from app.models.sale import Sale
    from app.models.order import Order
    before = (cu8.db.query(Sale).count(), cu8.db.query(Order).count())
    cu8.enabled['ml'] = False
    for path in ['predict-demand', 'customer-clusters']:
        result = cu8.client.post('/api/reports/ml/' + path, json={})
        assert result.status_code == 403, result.text
        assert 'plan' in result.text.lower()
    cu8.enabled['ml'] = True
    cu8.payload['rol'] = 'encargado'
    assert cu8.client.post('/api/reports/ml/predict-demand', json={}).status_code == 403
    cu8.payload['rol'] = 'administrador'
    result = cu8.client.post('/api/reports/ml/predict-demand', json={'branch_id': str(cu8.branches[1].id)})
    assert result.status_code == 404, result.text
    assert cu8.client.post('/api/reports/ml/predict-demand', json={'target_period_weeks': 13}).status_code == 422
    assert before == (cu8.db.query(Sale).count(), cu8.db.query(Order).count())


@pytest.mark.parametrize('body', [{}, None, {'use_synthetic_data': False}])
def test_ml_real_source_uses_only_own_sales_and_customer_not_cashier(cu8, body):
    forecast = cu8.client.post('/api/reports/ml/predict-demand', json={'target_period_weeks': 2})
    assert forecast.status_code == 200, forecast.text
    data = forecast.json()['data']
    assert len(data['forecast']) == 2
    assert data['data_source']['real_samples'] == 1
    from datetime import date
    from app.services.ml.demand_forecast_service import forecast_from_history
    day = date.today() - timedelta(days=14)
    week = day - timedelta(days=day.weekday())
    assert data == forecast_from_history({week: 3}, 2, today=date.today())
    result = cu8.client.post('/api/reports/ml/customer-clusters', **({'json': body} if body is not None else {}))
    assert result.status_code == 200, result.text
    data = result.json()['data']
    assert data['data_source']['real_samples'] == 1
    assert sum(segment['real_customers'] for segment in data['segments']) == 1
    from app.services.ml.customer_segmentation_service import segment_customers
    assert data == segment_customers([(14, 1, 25.0)])


@pytest.mark.parametrize('path', ['predict-demand', 'customer-clusters'])
def test_explicit_synthetic_http_trains_without_business_queries_or_writes(cu8, monkeypatch, path):
    from math import isfinite
    from sqlalchemy import select
    from app.core.database import Base

    tables = list(cu8.db.bind.dialect.get_table_names(cu8.db.connection()))
    business = [Base.metadata.tables[name] for name in tables]
    before = [cu8.db.execute(select(table)).all() for table in business]

    def forbidden(*args, **kwargs):
        pytest.fail('Synthetic ML must not query business data')

    with monkeypatch.context() as guarded:
        guarded.setattr(cu8.db, 'query', forbidden)
        guarded.setattr(cu8.db, 'execute', forbidden)
        results = [cu8.client.post('/api/reports/ml/' + path, json={'use_synthetic_data': True}) for _ in range(2)]
        for result in results:
            assert result.status_code == 200, result.text
        assert results[0].json() == results[1].json()
        data = results[0].json()['data']
        assert data['data_source']['mode'] == 'synthetic'
        assert data['data_source']['real_samples'] == 0
        assert data['data_source']['synthetic_samples'] > 0
        assert all(isfinite(value) for value in data['evaluation_metrics'].values())
        if path == 'predict-demand':
            assert len(data['forecast']) == 4
            assert data['evaluation_metrics']['test_samples'] > 0
            assert data['evaluation_metrics']['mse'] >= 0
            assert all(0 <= point['confidence_lower'] <= point['predicted_units'] <= point['confidence_upper'] for point in data['forecast'])
        else:
            assert data['clusters_count'] == 3
            assert -1 <= data['evaluation_metrics']['silhouette_score'] <= 1
            assert sum(s['size'] for s in data['segments']) == data['data_source']['synthetic_samples']
            assert all(s['real_customers'] == 0 for s in data['segments'])
        for enabled, role, detail in [(False, 'administrador', 'plan'), (True, 'encargado', None)]:
            cu8.enabled['ml'] = enabled
            cu8.payload['rol'] = role
            denied = cu8.client.post('/api/reports/ml/' + path, json={'use_synthetic_data': True})
            assert denied.status_code == 403, denied.text
            if detail:
                assert detail in denied.text.lower()
    assert [cu8.db.execute(select(table)).all() for table in business] == before


@pytest.mark.parametrize('kind', ['top_products', 'branch_comparison', 'critical_inventory'])
def test_nlp_executes_new_reports_and_preserves_query_contract(cu8, monkeypatch, kind):
    from app.services import natural_report_service
    from app.schemas.report_schema import NaturalReportInterpretation
    monkeypatch.setattr(natural_report_service.settings, 'gemini_api_key', 'cu8-test-stub')
    monkeypatch.setattr(natural_report_service, '_interpret_sync', lambda prompt: NaturalReportInterpretation(
        report_type=kind, interpretation='Consulta CU8 interpretada'))
    result = cu8.client.post('/api/reports/query', json={'query': 'Reporte de mi empresa'})
    assert result.status_code == 200, result.text
    data = result.json()['data']
    assert data['report_type'] == kind
    assert data['format'] == 'pdf'
    assert len(data['report']['rows']) == 1
    assert 'BetaCU8' not in result.text
    assert data['columns']


def test_manager_scope_is_preserved_for_new_reports_and_exports(cu8):
    cu8.payload['rol'] = 'encargado'
    cu8.payload['branch_id'] = str(cu8.branches[0].id)
    result = cu8.client.get('/api/reports/branch-comparison', params={'branch_id': str(cu8.branches[1].id)})
    assert result.status_code == 200, result.text
    assert [row['branch_id'] for row in result.json()['data']['rows']] == [str(cu8.branches[0].id)]
    result = cu8.client.get('/api/reports/export/excel', params={'type': 'branch_comparison', 'branch_id': str(cu8.branches[1].id)})
    assert result.status_code == 200
    values = list(load_workbook(BytesIO(result.content)).active.values)
    assert values[1] == ('AlphaCU8', 3, 75)
    assert 'BetaCU8' not in str(values)


def test_invalid_export_type_dates_and_missing_branch_do_not_mutate_data(cu8):
    from sqlalchemy import select
    from app.core.database import Base
    tables = [table for table in Base.metadata.sorted_tables if table.name in
              ['sales', 'sale_items', 'orders', 'order_items', 'inventory']]
    before = [cu8.db.execute(select(table)).all() for table in tables]
    assert cu8.client.get('/api/reports/export/excel', params={'type': 'sql'}).status_code == 422
    result = cu8.client.get('/api/reports/export/pdf', params={'from_date': '2026-05-02', 'to_date': '2026-05-01'})
    assert result.status_code == 422
    assert result.json()['detail'] == 'El rango de fechas no es válido'
    empty = cu8.client.get('/api/reports/top-products', params={'branch_id': str(uuid4())})
    assert empty.status_code == 200
    assert empty.json()['data']['rows'] == []
    assert [cu8.db.execute(select(table)).all() for table in tables] == before
