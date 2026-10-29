from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook


def test_excel_has_numeric_cells_literal_formulas_and_headers():
    from app.services.report_export_service import to_excel
    content = to_excel([{'product_name': '=HYPERLINK("bad")', 'gross_sales': Decimal('12.50')}], ['product_name', 'gross_sales'])
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.active
    assert sheet['A1'].value == 'Producto'
    assert sheet['A2'].value == '=HYPERLINK("bad")'
    assert sheet['A2'].data_type == 's'
    assert sheet['B2'].value == 12.5
    assert sheet.freeze_panes == 'A2'


@pytest.mark.parametrize('rows', [[], [{'product_name': '<Blusa> & "especial"', 'quantity_sold': 7}]])
def test_pdf_is_a_valid_document_even_for_empty_or_markup_text(rows):
    from app.services.report_export_service import to_pdf
    content = to_pdf('Ventas <2026> & sucursales', rows, ['product_name', 'quantity_sold'])
    assert content.startswith(b'%PDF-')
    assert content.rstrip().endswith(b'%%EOF')
    assert len(content) > 500


def test_empty_excel_keeps_column_headers():
    from app.services.report_export_service import to_excel
    sheet = load_workbook(BytesIO(to_excel([], ['branch_name', 'quantity_sold']))).active
    assert list(sheet.values) == [('Sucursal', 'Unidades vendidas')]


@pytest.mark.parametrize('kind', ['top_products', 'branch_comparison', 'critical_inventory'])
def test_natural_contract_supports_the_new_report_kinds(kind):
    from app.schemas.report_schema import NaturalReportInterpretation
    assert NaturalReportInterpretation(report_type=kind, interpretation='Reporte real').report_type == kind
