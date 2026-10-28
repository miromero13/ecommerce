"""Excel and paginated PDF exports for the existing scoped report rows."""
from enum import Enum
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

COLUMN_LABELS = {
    'branch_name': 'Sucursal', 'product_name': 'Producto', 'variant_sku': 'SKU', 'type': 'Tipo',
    'quantity_sold': 'Unidades vendidas', 'gross_sales': 'Ventas brutas',
    'payment_method': 'Método de pago', 'payment_status': 'Estado de pago', 'sale_status': 'Estado de venta',
    'quantity': 'Cantidad', 'reserved_quantity': 'Reservado', 'available_quantity': 'Disponible',
    'minimum_stock': 'Stock mínimo', 'movement_type': 'Movimiento', 'movements_count': 'Movimientos',
}


def _value(value):
    if isinstance(value, Enum):
        return value.value
    return '' if value is None else value


def to_excel(rows: list[dict], headers: list[str]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Reporte'
    sheet.append([COLUMN_LABELS.get(key, key) for key in headers])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='166534')
    for row in rows:
        sheet.append([_value(row.get(key)) for key in headers])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                # Untrusted names/SKUs are text, never executable spreadsheet formulas.
                cell.data_type = 's'
    for index, key in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(40, max(18, len(COLUMN_LABELS.get(key, key)) + 4))
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def to_pdf(title: str, rows: list[dict], headers: list[str]) -> bytes:
    output = BytesIO()
    styles = getSampleStyleSheet()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=24, rightMargin=24,
                                 topMargin=24, bottomMargin=24)
    cell_style = styles['BodyText']
    cell_style.fontSize = 8
    cell_style.leading = 10
    cells = [[Paragraph(escape(COLUMN_LABELS.get(key, key)), cell_style) for key in headers]]
    cells.extend([[Paragraph(escape(str(_value(row.get(key)))), cell_style) for key in headers] for row in rows])
    table = LongTable(cells, repeatRows=1, colWidths=[document.width / len(headers)] * len(headers))
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#dcfce7')),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.lightgrey), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
    ]))
    document.build([Paragraph(escape(title), styles['Title']), Spacer(1, 12), table])
    return output.getvalue()
